"""Authenticated SSE query endpoints with resumable Redis-backed streams."""

from __future__ import annotations

import asyncio
import json
import tempfile
import threading
import time
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time, timezone
from decimal import Decimal
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.background import BackgroundTask

from api.deps import get_async_session, get_current_user, get_redis
from auth.dependencies import get_current_user_from_header_or_query, UserContext
from auth.rbac import get_user_allowed_tables
from config import get_settings
from graph.workflow import run_workflow_stream_with_session
from services.query_runs import (
    HEARTBEAT_INTERVAL_SECONDS,
    PARTIAL_FLUSH_CHARS,
    PARTIAL_FLUSH_SECONDS,
    QUERY_CAPACITY_DUPLICATE,
    QUERY_CAPACITY_OK,
    TERMINAL_EVENT_TYPES,
    append_event,
    acquire_query_capacity,
    create_run,
    delete_run,
    get_run,
    is_cancel_requested_sync,
    release_query_capacity,
    request_cancel,
    set_run_message_id,
    stream_events,
    touch_heartbeat,
)
from utils.workflow_logger import WorkflowLogger, LogContext

router = APIRouter(tags=["query"])

_GEN_SENTINEL = object()
_BACKGROUND_TASKS: set[asyncio.Task] = set()


@dataclass
class RunControl:
    user_id: str
    session_id: str
    message_id: str | None
    cancel_event: threading.Event


class CancelQueryRequest(BaseModel):
    run_id: str


_RUN_CONTROLS: dict[str, RunControl] = {}


def _next_or_sentinel(sync_gen):
    try:
        return next(sync_gen)
    except StopIteration:
        return _GEN_SENTINEL


async def _aiter_sync_gen(sync_gen):
    while True:
        item = await asyncio.to_thread(_next_or_sentinel, sync_gen)
        if item is _GEN_SENTINEL:
            break
        yield item


def _format_history(messages: list) -> str:
    if not messages:
        return ""
    lines = ["历史对话:"]
    for m in messages:
        role = "用户" if m.role == "user" else "助手"
        lines.append(f"{role}: {m.content}")
    return "\n".join(lines)


def _sse_event(payload: dict, event_id: str | None = None) -> str:
    data = json.dumps(payload, ensure_ascii=False, default=str)
    if event_id:
        return f"id: {event_id}\ndata: {data}\n\n"
    return f"data: {data}\n\n"


def _single_error_stream(message: str, **extra: object) -> StreamingResponse:
    async def generate():
        yield _sse_event({"type": "error", "content": message, **extra})

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "Connection": "keep-alive"},
    )


def _stream_response(redis, run_id: str, from_event_id: str | None = None) -> StreamingResponse:
    async def generate():
        async for event_id, payload in stream_events(redis, run_id, from_event_id or "0-0"):
            yield _sse_event(payload, event_id)
            if payload.get("type") in TERMINAL_EVENT_TYPES:
                break

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "Connection": "keep-alive"},
    )


@router.post("/api/query/cancel")
async def cancel_query(
    body: CancelQueryRequest,
    user: UserContext = Depends(get_current_user),
):
    redis = get_redis()
    run = await get_run(redis, body.run_id)
    control = _RUN_CONTROLS.get(body.run_id)
    owner_id = run.get("user_id") if run else (control.user_id if control else None)
    if not owner_id or owner_id != user.user_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Run not found")
    await request_cancel(redis, body.run_id, "user_cancelled")
    if control:
        control.cancel_event.set()
    return {"status": "cancelling", "run_id": body.run_id}


@router.get("/api/query/resume")
async def resume_query_stream(
    run_id: str = Query(..., min_length=1),
    from_event_id: str = Query(default="0-0"),
    user: UserContext = Depends(get_current_user_from_header_or_query),
):
    redis = get_redis()
    run = await get_run(redis, run_id)
    if not run or run.get("user_id") != user.user_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Run not found")
    return _stream_response(redis, run_id, from_event_id or "0-0")


class ExportTooLargeError(Exception):
    """Raised when an export would exceed the configured row cap."""


def _validate_export_sql(sql: str, allowed_tables: list[str] | None) -> None:
    from graph.nodes import check_query_node

    check = check_query_node({
        "generated_sql": sql,
        "allowed_tables": allowed_tables,
        "log_context": None,
    })
    if not check.get("sql_valid"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=check.get("sql_check_message") or "SQL is not allowed for export.",
        )


def _xlsx_cell_value(value):
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool, datetime, date, datetime_time)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _write_sql_export_xlsx_sync(sql: str, result_id: str) -> tuple[Path, int]:
    from db.connection import get_engine
    from openpyxl import Workbook

    settings = get_settings()
    chunk_size = max(1, int(settings.export_chunk_size or 2000))
    max_rows = max(1, int(settings.export_max_rows or 1048575))
    timeout_ms = max(1, int(settings.export_statement_timeout_ms or 300000))

    temp = tempfile.NamedTemporaryFile(
        prefix=f"query_result_{result_id}_",
        suffix=".xlsx",
        delete=False,
    )
    temp_path = Path(temp.name)
    temp.close()

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("查询结果")
    row_count = 0

    try:
        engine = get_engine()
        with engine.connect() as conn:
            conn.execute(
                text("SELECT set_config('statement_timeout', :timeout_value, true)"),
                {"timeout_value": f"{timeout_ms}ms"},
            )
            stream_conn = conn.execution_options(stream_results=True)
            result = stream_conn.execute(text(sql))
            ws.append([str(column) for column in result.keys()])

            while True:
                rows = result.fetchmany(chunk_size)
                if not rows:
                    break
                if row_count + len(rows) > max_rows:
                    raise ExportTooLargeError(
                        f"导出结果超过 {max_rows} 行，请缩小查询范围后重试。"
                    )
                for row in rows:
                    ws.append([_xlsx_cell_value(cell) for cell in row])
                row_count += len(rows)

        wb.save(temp_path)
        return temp_path, row_count
    except Exception:
        try:
            ws.close()
        except Exception:
            pass
        try:
            wb.close()
        except Exception:
            pass
        try:
            temp_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise


@router.get("/api/query-results/{result_id}/export")
async def export_query_result(
    result_id: str,
    user: UserContext = Depends(get_current_user),
    db: AsyncSession = Depends(get_async_session),
):
    from db.crud.query_results import get_query_result_for_user

    result = await get_query_result_for_user(db, result_id, user.user_id)
    if not result:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Query result not found")
    if not result.sql:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Query result has no SQL to export")

    redis = get_redis()
    allowed_tables = await get_user_allowed_tables(user.user_id, redis, db)
    if allowed_tables == []:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="No data table access")
    _validate_export_sql(result.sql, allowed_tables)

    try:
        file_path, _row_count = await asyncio.to_thread(
            _write_sql_export_xlsx_sync,
            result.sql,
            result_id,
        )
    except ExportTooLargeError as exc:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail=str(exc)) from exc

    return FileResponse(
        file_path,
        filename=f"query_result_{result_id}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        background=BackgroundTask(lambda: file_path.unlink(missing_ok=True)),
    )


@router.get("/api/query")
async def query_stream(
    q: str = Query(..., description="User question"),
    session_id: str = Query(..., description="Session ID"),
    run_id: str = Query(..., min_length=1, description="Query run ID"),
    group_id: str = Query(default=None, description="Optional table group ID"),
    user: UserContext = Depends(get_current_user_from_header_or_query),
    db: AsyncSession = Depends(get_async_session),
):
    question = q.strip()
    if not question:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Question cannot be empty")
    if session_id == "default":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="A valid session ID is required")
    if run_id in _RUN_CONTROLS:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Run already exists")

    redis = get_redis()
    if await get_run(redis, run_id):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Run already exists")

    from db.crud.sessions import get_session

    session = await get_session(db, session_id)
    if not session or session.user_id != user.user_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")

    allowed_tables = await get_user_allowed_tables(user.user_id, redis, db)
    if allowed_tables == []:
        return _single_error_stream("The current user is not allowed to access any data tables.")

    settings = get_settings()
    conversation_history = ""
    try:
        from db.crud.messages import get_messages_by_session_sync

        messages = await asyncio.to_thread(
            get_messages_by_session_sync,
            session_id,
            settings.max_history_turns * 2,
        )
        conversation_history = _format_history(messages)
    except Exception as exc:
        print(f"[history] load failed: session_id={session_id}, error={exc}")

    available_results = []
    try:
        from db.crud.query_results import list_query_result_summaries_sync

        available_results = await asyncio.to_thread(
            list_query_result_summaries_sync,
            session_id,
            10,
        )
    except Exception as exc:
        print(f"[query_results] load failed: session_id={session_id}, error={exc}")

    group_table_filter = None
    if group_id:
        try:
            from db.crud.table_groups import get_group

            group = await get_group(db, group_id)
            if not group:
                return _single_error_stream("Selected data group does not exist.")
            group_table_names = [m.table_name for m in group.tables]
            if not group_table_names:
                return _single_error_stream("Selected data group has no bound tables.")
            if allowed_tables is not None:
                allowed_set = {t.lower() for t in allowed_tables}
                group_table_filter = [t for t in group_table_names if t.lower() in allowed_set]
            else:
                group_table_filter = group_table_names
            if not group_table_filter:
                return _single_error_stream("The current user cannot access tables in this group.")
        except Exception as exc:
            print(f"[query] table group load failed: group_id={group_id}, error={exc}")
            return _single_error_stream("Failed to read selected data group.")
    elif allowed_tables is None:
        try:
            from db.crud.table_groups import list_groups

            all_groups = await list_groups(db)
            all_table_set: set[str] = set()
            for group in all_groups:
                for member in group.tables:
                    all_table_set.add(member.table_name)
            if all_table_set:
                group_table_filter = list(all_table_set)
        except Exception as exc:
            print(f"[query] all-group table filter failed: error={exc}")

    capacity_status = await acquire_query_capacity(redis, run_id, user.user_id)
    if capacity_status == QUERY_CAPACITY_DUPLICATE:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Run already exists")
    if capacity_status != QUERY_CAPACITY_OK:
        return _single_error_stream(
            "当前查询人数较多，请稍后重试。",
            code="query_concurrency_limit",
            retry_after_seconds=5,
        )

    started_at = datetime.now(timezone.utc).isoformat()
    assistant_msg_id = None
    run_created = False
    try:
        from db.crud.messages import save_assistant_message_sync, save_user_message_sync

        created_run = await create_run(
            redis,
            run_id=run_id,
            user_id=user.user_id,
            session_id=session_id,
            message_id=None,
        )
        if not created_run:
            await release_query_capacity(redis, run_id, user.user_id)
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Run already exists")
        run_created = True

        await asyncio.to_thread(save_user_message_sync, session_id, question)
        assistant_msg_id = await asyncio.to_thread(
            save_assistant_message_sync,
            session_id,
            "",
            {
                "status": "streaming",
                "run_id": run_id,
                "started_at": started_at,
                "last_event_id": "0-0",
                "last_heartbeat": started_at,
            },
        )
        await set_run_message_id(redis, run_id, assistant_msg_id)
    except HTTPException:
        raise
    except Exception as exc:
        import traceback

        await release_query_capacity(redis, run_id, user.user_id)
        if run_created:
            await delete_run(redis, run_id)
        print(f"[messages] initialize failed: session_id={session_id}, error={exc}")
        print(f"[messages] traceback: {traceback.format_exc()}")
        return _single_error_stream("Query initialization failed.")

    cancel_event = threading.Event()
    _RUN_CONTROLS[run_id] = RunControl(
        user_id=user.user_id,
        session_id=session_id,
        message_id=assistant_msg_id,
        cancel_event=cancel_event,
    )

    if settings.log_enabled:
        WorkflowLogger.initialize(settings)
    log_context = LogContext("normal", question) if settings.log_enabled else None

    async def run_workflow_and_persist():
        start_time = time.time()
        final_state = None
        accumulated_answer = ""
        error_text = None
        stopped = False
        last_event_id = "0-0"
        last_flush_time = time.monotonic()
        last_flush_len = 0
        heartbeat_stop = asyncio.Event()

        async def heartbeat_loop():
            while not heartbeat_stop.is_set():
                try:
                    await asyncio.wait_for(heartbeat_stop.wait(), timeout=HEARTBEAT_INTERVAL_SECONDS)
                except asyncio.TimeoutError:
                    await touch_heartbeat(redis, run_id)

        async def flush_partial(force: bool = False):
            nonlocal last_flush_len, last_flush_time
            if not assistant_msg_id or not accumulated_answer:
                return
            now = time.monotonic()
            if (
                not force
                and now - last_flush_time < PARTIAL_FLUSH_SECONDS
                and len(accumulated_answer) - last_flush_len < PARTIAL_FLUSH_CHARS
            ):
                return

            from db.crud.messages import update_assistant_message_partial_sync

            await asyncio.to_thread(
                update_assistant_message_partial_sync,
                assistant_msg_id,
                accumulated_answer,
                {
                    "status": "streaming",
                    "run_id": run_id,
                    "started_at": started_at,
                    "last_event_id": last_event_id,
                    "last_heartbeat": datetime.now(timezone.utc).isoformat(),
                },
            )
            last_flush_time = now
            last_flush_len = len(accumulated_answer)

        def should_cancel() -> bool:
            if cancel_event.is_set():
                return True
            if is_cancel_requested_sync(run_id):
                cancel_event.set()
                return True
            return False

        heartbeat_task = asyncio.create_task(heartbeat_loop())
        try:
            workflow_gen = run_workflow_stream_with_session(
                question=question,
                session_id=session_id,
                user_id=user.user_id,
                allowed_tables=allowed_tables,
                log_context=log_context,
                conversation_history=conversation_history or None,
                group_table_filter=group_table_filter,
                available_results=available_results,
                should_cancel=should_cancel,
            )

            async for output in _aiter_sync_gen(workflow_gen):
                output_type = output.get("type")
                if output_type == "answer_chunk":
                    accumulated_answer += output.get("content", "")
                    last_event_id = await append_event(redis, run_id, output, progress=True)
                    await flush_partial()
                    continue

                if output_type == "done":
                    final_state = output.get("state", {}) or {}
                    execution_result = final_state.get("execution_result")
                    if execution_result:
                        last_event_id = await append_event(
                            redis,
                            run_id,
                            {"type": "result", "data": execution_result},
                            progress=True,
                        )
                    continue

                if output_type == "stopped":
                    stopped = True
                    final_state = output.get("state", {}) or {}
                    break

                last_event_id = await append_event(
                    redis,
                    run_id,
                    output,
                    progress=output_type in {"status", "explanation", "result"},
                )
        except Exception as exc:
            import traceback

            error_text = str(exc)
            print(f"[workflow] failed: {exc}\n{traceback.format_exc()}")
        finally:
            heartbeat_stop.set()
            heartbeat_task.cancel()
            try:
                await heartbeat_task
            except asyncio.CancelledError:
                pass
            await flush_partial(force=True)

            message_status = "stopped" if stopped or cancel_event.is_set() else ("failed" if error_text else "completed")
            final_answer = (
                (final_state or {}).get("final_answer")
                or accumulated_answer
                or (f"Query failed: {error_text}" if error_text else "")
                or ""
            )
            query_result_id = None

            if (
                final_state
                and message_status == "completed"
                and final_state.get("execution_success")
                and final_state.get("execution_result")
            ):
                try:
                    from auth.rbac import extract_tables_from_sql
                    from db.crud.query_results import build_query_result_summary, save_query_result_sync

                    generated_sql = final_state.get("generated_sql") or ""
                    execution_result = final_state.get("execution_result") or {}
                    referenced_tables = extract_tables_from_sql(generated_sql) if generated_sql else []
                    query_result_id = await asyncio.to_thread(
                        save_query_result_sync,
                        session_id,
                        assistant_msg_id,
                        question,
                        final_state.get("fused_question"),
                        generated_sql,
                        execution_result,
                        build_query_result_summary(question, execution_result),
                        referenced_tables,
                    )
                    final_state["query_result_id"] = query_result_id
                except Exception as exc:
                    print(f"[query_results] save failed: message_id={assistant_msg_id}, error={exc}")

            if message_status == "stopped":
                terminal_payload = {"type": "stopped", "state": final_state or {}}
            elif message_status == "failed":
                terminal_payload = {"type": "error", "content": error_text or "Query failed."}
            else:
                terminal_payload = {"type": "done", "state": final_state or {}}
            last_event_id = await append_event(redis, run_id, terminal_payload)

            try:
                from db.crud.messages import update_assistant_message_if_streaming_sync

                metadata = {
                    "sql": (final_state or {}).get("generated_sql"),
                    "explanation": (final_state or {}).get("query_explanation"),
                    "result": (final_state or {}).get("execution_result"),
                    "status": message_status,
                    "run_id": run_id,
                    "started_at": started_at,
                    "finished_at": datetime.now(timezone.utc).isoformat(),
                    "last_event_id": last_event_id,
                    "intent_type": (final_state or {}).get("intent_type"),
                    "referenced_result_ids": (final_state or {}).get("referenced_result_ids", []),
                    "query_result_id": query_result_id,
                }
                if message_status == "stopped":
                    metadata["stopped_at"] = datetime.now(timezone.utc).isoformat()
                if error_text:
                    metadata["error"] = error_text
                await asyncio.to_thread(
                    update_assistant_message_if_streaming_sync,
                    assistant_msg_id,
                    final_answer,
                    metadata,
                )
            except Exception as exc:
                import traceback

                print(f"[messages] assistant update failed: message_id={assistant_msg_id}, error={exc}")
                print(f"[messages] traceback: {traceback.format_exc()}")

            if final_state:
                try:
                    from tasks.audit import write_audit_log_task

                    elapsed_ms = int((time.time() - start_time) * 1000)
                    result = final_state.get("execution_result") or {}
                    await asyncio.to_thread(
                        write_audit_log_task.delay,
                        user_id=user.user_id,
                        session_id=session_id,
                        question=question,
                        generated_sql=final_state.get("generated_sql", ""),
                        execution_success=final_state.get("execution_success", False),
                        execution_time_ms=elapsed_ms,
                        row_count=result.get("row_count"),
                        error_message=final_state.get("error_message"),
                    )
                except Exception:
                    pass

            _RUN_CONTROLS.pop(run_id, None)
            try:
                await release_query_capacity(redis, run_id, user.user_id)
            except Exception:
                pass

    task = asyncio.create_task(run_workflow_and_persist())
    _BACKGROUND_TASKS.add(task)
    task.add_done_callback(_BACKGROUND_TASKS.discard)

    return _stream_response(redis, run_id, "0-0")
